"""Exporta o(s) quadro(s) de Project Review para um .xlsx formatado como o dashboard."""

import datetime as dt
import re

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side

PEACH = PatternFill("solid", fgColor="F4C7A1")
PEACH_LIGHT = PatternFill("solid", fgColor="FADCC4")
PEACH_PALE = PatternFill("solid", fgColor="FDF0E6")
BLUE = PatternFill("solid", fgColor="CFDAF0")
GRID = Side(style="thin", color="B7B7B7")
DOTTED = Side(style="dotted", color="7F7F7F")
BOLD = Font(bold=True)
CENTER = Alignment(horizontal="center", vertical="center", wrap_text=True)
RIGHT = Alignment(horizontal="right", vertical="center")
NUM_FMT = '#,##0;-#,##0;0'
PCT_FMT = '0.00%'
DATE_FMT = 'dd/mm/yyyy'


def _sheet_title(text, used):
    t = re.sub(r"[\\/*?:\[\]]", "", text)[:31] or "Project"
    base, i = t, 2
    while t in used:
        suffix = f" ({i})"
        t = base[: 31 - len(suffix)] + suffix
        i += 1
    used.add(t)
    return t


def write_project(ws, group, p):
    ncol = len(p["columns"])
    last = 2 + ncol  # coluna do Δ (1-based: A=labels)
    ws.column_dimensions["A"].width = 34
    for c in range(2, last + 1):
        ws.column_dimensions[ws.cell(1, c).column_letter].width = 20
    ws.column_dimensions[ws.cell(1, last).column_letter].width = 10

    ws.cell(1, 1, f"{group} — {p['sheet']}").font = Font(italic=True, color="7F7F7F")

    r = 3
    ws.cell(r, 1, p["name"])
    for i, col in enumerate(p["columns"]):
        ws.cell(r, 2 + i, "\n".join(col["lines"]))
    ws.cell(r, last, "Δ")
    for c in range(1, last + 1):
        cell = ws.cell(r, c)
        cell.font, cell.alignment = BOLD, CENTER
        cell.fill = PEACH_LIGHT if c == last - 1 else PEACH
        cell.border = Border(top=GRID, bottom=GRID, left=GRID, right=GRID)
    ws.row_dimensions[r].height = 48

    r += 1
    for row in p["rows"]:
        r += 1
        fill = BLUE if row["kind"] == "orion" else PEACH if row["kind"] == "total" else None
        ws.cell(r, 1, row["label"])
        for i, v in enumerate(row["values"] + [row["delta"]]):
            ws.cell(r, 2 + i, v)
        for c in range(1, last + 1):
            cell = ws.cell(r, c)
            if c > 1:
                cell.alignment = RIGHT
                cell.number_format = PCT_FMT if row["percent"] else NUM_FMT
            if fill:
                cell.fill, cell.font = fill, BOLD
            elif c == last - 1:
                cell.fill = PEACH_PALE
            cell.border = Border(left=GRID, right=GRID, bottom=DOTTED)

    if p["footnote"] or p["notes"]:
        r += 2
        if p["footnote"]:
            ws.cell(r, 1, p["footnote"]).font = Font(italic=True, size=9)
        for n in p["notes"]:
            ws.cell(r, last - 2 if ncol >= 2 else 2, n["label"]).font = Font(size=9)
            v = ws.cell(r, last - 1, n["value"])
            v.number_format, v.font = PCT_FMT, Font(size=9)
            r += 1

    if p["kpis"]:
        r += 2
        ws.cell(r, 1, p["kpi_title"] or "KPIs").font = BOLD
        for k in p["kpis"]:
            r += 1
            ws.cell(r, 1, k["label"])
            for i, lines in enumerate(k["values"]):
                ws.cell(r, 2 + i, "\n".join(lines))
            for c in range(1, last):
                cell = ws.cell(r, c)
                cell.fill = BLUE
                cell.font = BOLD if c == 1 else Font()
                cell.alignment = CENTER if c > 1 else Alignment(vertical="center")
                cell.border = Border(bottom=DOTTED)
            ws.row_dimensions[r].height = 32

    if p.get("financing"):
        write_financing(ws, r + 2, p["financing"], last)

    ws.sheet_view.showGridLines = False


def _date(iso):
    try:
        return dt.date.fromisoformat(iso)
    except (TypeError, ValueError):
        return iso


def write_financing(ws, r, f, last):
    """Bloco "Financing" (dados de financing.json) a partir da linha r."""
    ws.cell(r, 1, "Financing" + (f" — {f['facility']}" if f.get("facility") else "")).font = BOLD
    for c in range(1, last + 1):
        ws.cell(r, c).border = Border(top=Side(style="medium", color="F26641"))

    def term(months, end):
        if not months:
            return None
        return f"{months} months" + (f" (until {_date(end):%d/%m/%Y})" if end else "")

    rate = " + ".join(x for x in [f.get("index"), f"{f['spread']:.3f}%".replace(".", ",")
                                  if isinstance(f.get("spread"), (int, float)) else None] if x)
    items = [
        ("Bank", f.get("bank"), None),
        ("Borrower", f.get("borrower"), None),
        ("Signed", _date(f.get("signed")), DATE_FMT),
        ("Maturity", _date(f.get("maturity")), DATE_FMT),
        ("Amount (€)", f.get("amount"), NUM_FMT),
        ("Term", term(f.get("term_months"), f.get("maturity")), None),
        ("Availability period", term(f.get("availability_months"), f.get("availability_end")), None),
        ("Interest rate", rate or None, None),
        ("Purpose", f.get("purpose"), None),
        ("Own funds required (€)", f.get("own_funds"), NUM_FMT),
        ("", f.get("own_funds_note"), None),
    ]
    for label, value, fmt in items:
        if value in (None, ""):
            continue
        r += 1
        ws.cell(r, 1, label).font = Font(color="7F7F7F")
        v = ws.cell(r, 2, value)
        v.alignment = Alignment(horizontal="left", vertical="center")
        if fmt:
            v.number_format = fmt

    d = f.get("distributions")
    if d:
        r += 2
        title = "Distributions to promoter"
        if isinstance(d.get("max"), (int, float)):
            title += f" — up to {d['max']:,.0f} €".replace(",", " ")
        ws.cell(r, 1, title).font = BOLD
        for text in d.get("conditions", []) + ([d["note"]] if d.get("note") else []):
            r += 1
            cell = ws.cell(r, 1, ("• " if text != d.get("note") else "") + text)
            cell.alignment = Alignment(wrap_text=True, vertical="top")
            ws.merge_cells(start_row=r, start_column=1, end_row=r, end_column=last)
            ws.row_dimensions[r].height = 30


ORANGE = PatternFill("solid", fgColor="F26641")
DARK = PatternFill("solid", fgColor="333333")
WHITE_BOLD = Font(bold=True, color="FFFFFF")


def write_budget(ws, group, p):
    """Budget detalhado: rubricas com as subrubricas agrupadas (outline) por baixo."""
    b = p["budget"]
    cols = b["columns"]
    last = 1 + len(cols)
    ws.column_dimensions["A"].width = 46
    for c in range(2, last + 1):
        ws.column_dimensions[ws.cell(1, c).column_letter].width = 16
    ws.cell(1, 1, f"{group} — {b['sheet']}").font = Font(italic=True, color="7F7F7F")

    # Cabeçalho: faixa de grupos + títulos das colunas (+ fórmulas tipo "(5)=(3)+(4)")
    ws.cell(3, 1, b.get("title") or p["name"])
    prev = None
    for i, col in enumerate(cols):
        c = 2 + i
        if col["group"] != prev:
            ws.cell(3, c, col["group"])
            prev = col["group"]
        ws.cell(4, c, "\n".join(col["lines"]))
        if col.get("sub"):
            ws.cell(5, c, col["sub"])
    for r in (3, 4, 5):
        for c in range(1, last + 1):
            cell = ws.cell(r, c)
            g = cols[c - 2]["group"].upper() if c > 1 else "BUDGET"
            cell.fill = PEACH if g.startswith("COST") else DARK if g.startswith(("VARIATION", "TOTAL INV")) else ORANGE
            cell.font = BOLD if g.startswith("COST") else WHITE_BOLD
            cell.alignment = CENTER if c > 1 else Alignment(vertical="bottom", wrap_text=True)
            cell.border = Border(left=GRID, right=GRID)
    ws.row_dimensions[4].height = 44
    ws.freeze_panes = "B6"

    r = 5
    def put(row, fill=None, font=None, level=0):
        nonlocal r
        r += 1
        label = row["label"]
        if row.get("code") not in (None, ""):
            code = row["code"]
            label = f"{code:g}  {label}" if isinstance(code, float) else f"{code}  {label}"
        ws.cell(r, 1, ("    " if level else "") + label)
        for i, v in enumerate(row["values"]):
            cell = ws.cell(r, 2 + i, v)
            cell.alignment = RIGHT
            cell.number_format = PCT_FMT if (row.get("percent") or cols[i]["percent"]) else NUM_FMT
        for c in range(1, last + 1):
            cell = ws.cell(r, c)
            if fill:
                cell.fill = fill
            if font:
                cell.font = font
            cell.border = Border(left=GRID, right=GRID, bottom=DOTTED)
        if level:
            ws.row_dimensions[r].outlineLevel = 1
            ws.row_dimensions[r].hidden = True

    for row in b["rows"]:
        if row["kind"] == "rubric":
            put(row, PEACH_PALE, BOLD)
            for child in row["children"]:
                put(child, font=Font(color="555555"), level=1)
        elif row["kind"] in ("total", "orion"):
            put(row, BLUE if row["kind"] == "orion" else PEACH_LIGHT, BOLD)
        else:
            put(row, font=BOLD)
    ws.sheet_properties.outlinePr.summaryBelow = False
    ws.sheet_view.showGridLines = False


def build_workbook(projects):
    """projects: lista de (grupo/ficheiro, projeto). Com budget detalhado, acrescenta uma folha "… Budget"."""
    wb = Workbook()
    wb.remove(wb.active)
    used = set()
    for group, p in projects:
        name = p.get("label") or p["name"]
        ws = wb.create_sheet(_sheet_title(name, used))
        write_project(ws, group, p)
        if p.get("budget"):
            short = name
            while len(short) > 24 and " " in short:  # corta numa palavra inteira (máx. 31 caracteres)
                short = short.rsplit(" ", 1)[0].rstrip(" -|")
            write_budget(wb.create_sheet(_sheet_title(f"{short} Budget", used)), group, p)
    return wb
