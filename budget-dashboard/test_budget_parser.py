"""Testes do parser com um workbook sintético (sem dados reais).  python -m unittest"""

import datetime as dt
import os
import re
import tempfile
import unittest
import zipfile

from openpyxl import Workbook

from budget_parser import apply_versions, attach_budget_details, parse_file, parse_workbook
from excel_export import build_workbook

LINES = ["1- Land costs", "2- Charges", "4- Construction", "12-Other Adjustments"]


def make_review_sheet(ws, name):
    ws["B1"], ws["C1"] = "Last Project Review Date", dt.datetime(2026, 7, 16)
    ws["B2"], ws["C2"] = "Next project review date", dt.datetime(2026, 10, 21)
    ws["B6"] = name
    ws["C6"] = "Investment Committee \nDec/24"
    ws["D6"] = "Project Review        16 / 07 / 2026        Orion"
    ws["E6"] = "Project Review        21 / 10 / 2026        Orion"
    ws["F6"], ws["H6"] = "Δ", "Residential "
    ws["I6"], ws["H7"], ws["I7"] = "Total", 7779.2, 8957.15
    ws["H9"], ws["I9"] = 204.41, 204.41
    r = 9
    for i, label in enumerate(LINES):
        ws.cell(r, 2, label)
        for c in range(3, 6):
            ws.cell(r, c, 100.0 * (i + 1) + c)
        ws.cell(r, 6, 1.0)
        r += 1
    ws.cell(20, 2, "TOTAL COST (PROJECT)"); ws.cell(20, 5, 1000.0); ws.cell(20, 6, 5.0)
    ws.cell(27, 2, "MARGIN (€)*"); ws.cell(27, 5, 200.0); ws.cell(27, 6, 0.0)
    ws.cell(28, 2, "MARGIN (ORIONS VIEW) (€)*"); ws.cell(28, 5, 190.0); ws.cell(28, 6, 0.0)
    ws.cell(29, 2, "MARGIN (%)*"); ws.cell(29, 5, 0.2359); ws.cell(29, 6, 0.0019)
    ws.cell(31, 2, "*Margin Pre-Tax")
    ws.cell(31, 4, "Margin w/out internal fees"); ws.cell(31, 5, 0.2627)
    ws.cell(33, 2, "Orions View KPIs (IRR / EM / Profit / Equity)")
    ws.cell(34, 2, "Unlevered")
    for c in range(3, 6):
        ws.cell(34, c, "16,5% / 1,3x \n 12.0M€ / -42.4M€")
    ws.cell(38, 5, 0.17); ws.cell(38, 6, "Margin post tax")


def make_budget_sheet(ws):
    ws["A3"], ws["B3"], ws["C3"] = "Code", "TEST\n(€, taxes included)", "BUDGET € + VAT (23%)"
    ws["F3"], ws["I3"], ws["J3"] = "COST OF THE PROJECT € + VAT", "VARIATION", "COMMENTS (required for each variation)"
    ws["C4"], ws["D4"], ws["E4"] = "Investment Comittee", "Project Review 01/01/2026", "Project Review 16/07/2026 Orion"
    ws["F4"], ws["G4"], ws["H4"] = "Signed commitments", "Forecasted commitments", "TOTAL"
    ws["F5"], ws["H5"], ws["I5"] = "(3)", "(5)=(3)+(4)", "(6)=(5)-(2)"
    rows = [
        (None, "1- Land costs / Terrain", 100), (111, "Plot / Terreno", 90), (112, "Notary fees", 10),
        (None, "2- Charges", 20), (151, "License fees", 20),
        (None, "TOTAL COST", 120), (None, "MARGIN (%)", 0.25),
        (None, "Unlevered", None), (None, "IRR ANNUAL", 0.3),
    ]
    for i, (code, label, v) in enumerate(rows):
        r = 7 + i
        ws.cell(r, 1, code)
        ws.cell(r, 2, label)
        if v is not None:
            for c in range(3, 10):
                ws.cell(r, c, v)
        ws.cell(r, 10, "comentário")
    ws.column_dimensions["D"].hidden = True  # revisão antiga escondida no Excel
    for r in (8, 9):  # subrubricas agrupadas e recolhidas: continuam a ser lidas
        ws.row_dimensions[r].hidden = True


class ParserTest(unittest.TestCase):
    @classmethod
    def setUpClass(cls):
        wb = Workbook()
        wb.active.title = "Cover"
        wb.active["A1"] = "TOTAL COST"  # sem Δ: não é uma folha de review
        ws = wb.create_sheet("PR NOLA")
        make_review_sheet(ws, "NOLA")
        # segundo quadro na mesma folha (outra vista de custos)
        ws["B45"], ws["C45"], ws["D45"], ws["F45"] = "NOLA", "Acquisition Orion", "Project Review 21 / 10 / 2026", "Δ"
        ws["B48"], ws["C48"], ws["D48"], ws["F48"] = "1- Land costs", 19934.7, 686.5, 0
        ws["B59"], ws["D59"], ws["F59"] = "TOTAL COST (PROJECT)", 79220.7, 295.4
        ws["D74"], ws["E74"] = "Margin w/out internal fees", 0.2627
        make_review_sheet(wb.create_sheet("PR OTHER"), "OTHER")
        cmp_ws = wb.create_sheet("Comparison")  # tem Δ, mas "Total costs" não é TOTAL COST
        cmp_ws["B22"], cmp_ws["D22"], cmp_ws["F22"] = "Key milestones", "Current", "Δ"
        cmp_ws["B33"], cmp_ws["D33"], cmp_ws["F33"] = "Total costs (k€)", 66629.1, 1950.1
        hidden = wb.create_sheet("PR HIDDEN")
        make_review_sheet(hidden, "HIDDEN")
        hidden.sheet_state = "hidden"
        make_budget_sheet(wb.create_sheet("BUDGET FASE 1"))
        fd, cls.path = tempfile.mkstemp(suffix=".xlsx")
        os.close(fd)
        wb.save(cls.path)
        cls.projects = parse_workbook(cls.path)

    @classmethod
    def tearDownClass(cls):
        os.remove(cls.path)

    def test_detects_review_sheets_only(self):
        self.assertEqual([p["name"] for p in self.projects], ["NOLA", "OTHER"])
        self.assertEqual(len(parse_workbook(self.path, include_hidden=True)), 3)

    def test_header_and_dates(self):
        p = self.projects[0]
        self.assertEqual([c["lines"] for c in p["columns"]], [
            ["Investment Committee", "Dec/24"],
            ["Project Review", "16/07/2026", "Orion"],
            ["Project Review", "21/10/2026", "Orion"],
        ])
        self.assertEqual((p["last_review"], p["next_review"]), ("2026-07-16", "2026-10-21"))

    def test_rows(self):
        rows = {r["label"]: r for r in self.projects[0]["rows"]}
        self.assertEqual(list(rows)[:4], LINES)
        self.assertEqual(rows["4- Construction"]["values"], [303.0, 304.0, 305.0])
        self.assertEqual(rows["TOTAL COST (PROJECT)"]["kind"], "total")
        self.assertEqual(rows["MARGIN (ORIONS VIEW) (€)*"]["kind"], "orion")
        self.assertTrue(rows["MARGIN (%)*"]["percent"])
        self.assertAlmostEqual(rows["MARGIN (%)*"]["delta"], 0.0019)
        self.assertNotIn("*Margin Pre-Tax", rows)

    def test_second_table(self):
        p = self.projects[0]
        self.assertEqual(len(p["views"]), 1)
        v = p["views"][0]
        self.assertEqual([r["label"] for r in v["rows"]], ["1- Land costs", "TOTAL COST (PROJECT)"])
        self.assertEqual(v["rows"][0]["values"], [19934.7, 686.5, None])
        # as notas do segundo quadro não passam para o principal
        self.assertEqual(len([n for n in p["notes"] if n["label"].startswith("Margin w/out")]), 1)
        self.assertEqual(self.projects[1]["views"], [])

    def test_sqm_block(self):
        p = self.projects[0]
        self.assertEqual(p["sqm_columns"], [{"label": "Residential", "area": 7779.2}, {"label": "Total", "area": 8957.15}])
        rows = {r["label"]: r for r in p["rows"]}
        self.assertEqual(rows["1- Land costs"]["sqm"], [204.41, 204.41])
        self.assertNotIn("sqm", rows["2- Charges"])

    def test_notes_and_kpis(self):
        p = self.projects[0]
        self.assertEqual(p["footnote"], "*Margin Pre-Tax")
        self.assertEqual([(n["label"], n["value"], n["below_kpis"]) for n in p["notes"]],
                         [("Margin w/out internal fees", 0.2627, False), ("Margin post tax", 0.17, True)])
        self.assertEqual(p["kpis"][0]["label"], "Unlevered")
        self.assertEqual(p["kpis"][0]["values"][0], ["16,5% / 1,3x", "12.0M€ / -42.4M€"])

    def test_budget_detail(self):
        b = parse_file(self.path)["budgets"]["BUDGET FASE 1"]
        # Coluna D escondida é lida mas marcada; COMMENTS fica de fora.
        self.assertEqual([c["lines"][0] for c in b["columns"]],
                         ["Investment Comittee", "Project Review", "Project Review", "Signed commitments",
                          "Forecasted commitments", "TOTAL", "VARIATION"])
        self.assertEqual([c["hidden"] for c in b["columns"]], [False, True] + [False] * 5)
        self.assertEqual([c["group"] for c in b["columns"]][-1], "VARIATION")
        self.assertEqual(b["columns"][3]["sub"], "(3)")
        land = b["rows"][0]
        self.assertEqual((land["kind"], land["label"]), ("rubric", "1- Land costs / Terrain"))
        self.assertEqual([(c["code"], c["label"]) for c in land["children"]],
                         [(111.0, "Plot / Terreno"), (112.0, "Notary fees")])
        self.assertEqual([r["label"] for r in b["rows"]],
                         ["1- Land costs / Terrain", "2- Charges", "TOTAL COST", "MARGIN (%)"])  # pára nos KPIs
        self.assertTrue(b["rows"][-1]["percent"])

    def test_invalid_defined_names(self):
        # Print_Titles = #N/A faz o openpyxl recusar o ficheiro; a app ignora os nomes.
        fd, path = tempfile.mkstemp(suffix=".xlsx")
        os.close(fd)
        try:
            with zipfile.ZipFile(self.path) as zin, zipfile.ZipFile(path, "w") as zout:
                for item in zin.infolist():
                    data = zin.read(item.filename)
                    if item.filename == "xl/workbook.xml":
                        data = re.sub(rb"<definedNames>.*?</definedNames>|(?=<calcPr)",
                                      b'<definedNames><definedName name="_xlnm.Print_Titles" '
                                      b'localSheetId="1">#N/A</definedName></definedNames>',
                                      data, count=1, flags=re.S)
                    zout.writestr(item, data)
            self.assertEqual([p["name"] for p in parse_workbook(path)], ["NOLA", "OTHER"])
        finally:
            os.remove(path)

    def test_excel_export(self):
        wb = build_workbook([("File", p) for p in self.projects])
        self.assertEqual(wb.sheetnames, ["NOLA", "OTHER"])
        self.assertEqual(wb["NOLA"]["A3"].value, "NOLA")

    def test_attach_budget_details(self):
        # Quadro: IC | PR jul | PR out | Δ (= out − jul). Budget: colunas por outra ordem, uma escondida.
        rows = [("1- Land costs", [100, 110, 120], 10), ("2- Charges", [20, 25, 30], 5),
                ("4- Construction", [500, 520, 540], 20), ("TOTAL COST (PROJECT)", [620, 655, 690], 35)]
        project = {"name": "X", "columns": [{"lines": ["Investment Committee"]},
                                            {"lines": ["Project Review", "16/07/2026"]},
                                            {"lines": ["Project Review", "21/10/2026"]}],
                   "rows": [{"label": l, "kind": "total" if l.startswith("TOTAL") else "line",
                             "percent": False, "values": v, "delta": d} for l, v, d in rows]}
        col = lambda *lines, hidden=False: {"group": "BUDGET", "lines": list(lines), "hidden": hidden, "percent": False}
        # colunas do budget: [PR 2026-07-16 (escondida), TOTAL, IC, PR 16/07/2026 duplicada visível]
        budget = {"sheet": "BUDGET FASE 1", "columns": [col("Project Review 2026-07-16", hidden=True), col("TOTAL"),
                                                         col("Investment Comittee"), col("Project Review", "16/07/2026")],
                  "rows": [
                      {"label": "1- Land costs / Terrain", "kind": "rubric", "values": [110, 120, 100, 110], "children": [
                          {"label": "Plot", "code": 111.0, "values": [100, 105, 90, 100]},
                          {"label": "Notary", "code": 112.0, "values": [10, 15, 10, 10]}]},
                      {"label": "2- Charges", "kind": "rubric", "values": [25, 30, 20, 25], "children": []},
                      {"label": "4- Construction", "kind": "rubric", "values": [520, 540, 500, 520], "children": [
                          {"label": "Gal Contractor", "code": 321.1, "values": [520, 540, 500, 520]}]},
                  ]}
        q = attach_budget_details(project, budget)
        self.assertEqual(q["budget_link"]["delta"], "difference")
        # PR jul -> coluna visível com a mesma data (não a escondida); PR out -> TOTAL
        self.assertEqual(q["budget_link"]["columns"], [["Investment Comittee"], ["Project Review", "16/07/2026"], ["TOTAL"]])
        land = q["rows"][0]["children"]
        self.assertEqual([(c["code"], c["values"], c["delta"]) for c in land],
                         [("111", [90, 100, 105], 5), ("112", [10, 10, 15], 5)])
        self.assertNotIn("children", q["rows"][1])  # rubrica sem subrubricas
        self.assertEqual(q["rows"][2]["children"][0]["code"], "321.1")
        self.assertNotIn("children", project["rows"][0])  # o original não é alterado


        # Versões: colunas BUDGET (incl. escondida) e o TOTAL; linhas ligadas ao budget
        self.assertEqual([v["col"] for v in q["versions"]], [0, 1, 2, 3])  # aqui todas são do grupo BUDGET
        self.assertEqual(q["version_default"], [2, 3, 1])
        self.assertEqual(q["rows"][0]["bvalues"], [110, 120, 100, 110])
        self.assertNotIn("bvalues", q["rows"][3])  # TOTAL COST não existe na folha de budget de teste
        # Escolher a coluna escondida (0) como comparação e o TOTAL (1) como referência: Δ recalculado
        q2 = q
        v = apply_versions(q2, [0, 1])
        self.assertEqual(v["columns"][1]["lines"], ["Project Review 2026-07-16"])
        self.assertEqual(v["columns"][2]["lines"], ["Project Review", "21/10/2026"])  # referência = padrão
        self.assertEqual((v["rows"][0]["values"], v["rows"][0]["delta"]), ([100, 110, 120], 10))
        self.assertEqual([c["values"] for c in v["rows"][0]["children"]], [[90, 100, 105], [10, 10, 15]])
        self.assertEqual(v["rows"][3]["values"], [620, None, 690])  # sem dados no budget: vazio (Δ também)
        self.assertIsNone(v["rows"][3]["delta"])
        self.assertIs(apply_versions(q2, [3, 1]), q2)  # escolha = padrão: o quadro fica igual

        wb = build_workbook([("File", {**self.projects[0], **q, "footnote": None, "notes": [], "kpis": []})])
        ws = wb.active
        labels = [ws.cell(r, 1).value for r in range(5, 9)]
        self.assertEqual(labels, ["1- Land costs", "    111  Plot", "    112  Notary", "2- Charges"])
        self.assertEqual(ws.row_dimensions[6].outlineLevel, 1)
        self.assertTrue(ws.row_dimensions[6].hidden)


    def test_commissions_511(self):
        # rubrica 511 (comissões de mediadores): orçamento na coluna do último PR, adjudicado em "Signed commitments"
        project = {"name": "X", "columns": [{"lines": ["Investment Committee"]}, {"lines": ["Project Review", "21/10/2026"]}],
                   "rows": [{"label": l, "kind": "line", "percent": False, "values": v, "delta": None}
                            for l, v in [("1- Land costs", [100, 120]), ("2- Charges", [20, 30]), ("9-Marketing cost", [50, 60])]]}
        col = lambda group, *lines: {"group": group, "lines": list(lines), "hidden": False, "percent": False}
        budget = {"sheet": "BUDGET", "columns": [col("BUDGET", "Investment Comittee"), col("BUDGET", "Project Review 21/10/2026"),
                                                 col("COST OF THE PROJECT", "Signed commitments"), col("COST OF THE PROJECT", "TOTAL")],
                  "rows": [
                      {"label": "1- Land costs", "kind": "rubric", "values": [100, 120, 80, 120], "children": []},
                      {"label": "2- Charges", "kind": "rubric", "values": [20, 30, 10, 30], "children": []},
                      {"label": "9-Marketing cost", "kind": "rubric", "values": [50, 60, 25, 60], "children": [
                          {"label": "External sales fees", "code": 511.0, "values": [40, 45, 18, 45]}]},
                  ]}
        c = attach_budget_details(project, budget)["commissions"]
        self.assertEqual((c["label"], c["budget"], c["signed"], c["available"]), ("External sales fees", 45, 18, 27))
if __name__ == "__main__":
    unittest.main()
