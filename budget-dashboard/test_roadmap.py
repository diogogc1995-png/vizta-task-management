"""Testes do leitor do roadmap (folha sintética, sem dados reais)."""

import datetime as dt
import os
import tempfile
import unittest

from openpyxl import Workbook

import roadmap

D = lambda y, m: dt.datetime(y, m, 1)


def make_sheet(ws):
    ws["B4"] = "Project / Budget review"
    for col, text in {"J": "PSPA", "L": "Acquisition Date", "M": "Projeto Base", "N": "Licensing/Permit",
                      "P": "Commercial Launch (pre-sales)", "Q": "Construction", "S": "End of deliveries"}.items():
        ws[f"{col}5"] = text
    ws["C6"], ws["F6"], ws["N6"], ws["O6"], ws["Q6"], ws["R6"] = "BUILDING NAMES", "Nb. Apartments", "Start", "End", "Start", "End"
    rows = [
        ("Zona A", "Torre I", 100, None, D(2019, 3), D(2020, 7), D(2021, 4), D(2023, 4), D(2023, 9)),
        (None, "Torre II", 120, None, D(2019, 3), D(2022, 9), D(2023, 7), D(2026, 5), D(2026, 6)),
        ("Zona B", "A", 80, D(2023, 11), D(2024, 3), D(2025, 2), D(2025, 7), D(2027, 3), D(2027, 9)),
        ("Vila", None, 159, None, D(2022, 9), D(2022, 1), D(2022, 11), D(2026, 7), D(2026, 7)),
        (None, "Info changes since PR", None, None, None, None, None, None, None),
    ]
    for i, (area, name, apts, pspa, acq, launch, cs, ce, eod) in enumerate(rows, start=8):
        ws[f"B{i}"], ws[f"C{i}"], ws[f"F{i}"], ws[f"J{i}"], ws[f"L{i}"] = area, name, apts, pspa, acq
        ws[f"P{i}"], ws[f"Q{i}"], ws[f"R{i}"], ws[f"S{i}"] = launch, cs, ce, eod


class RoadmapTest(unittest.TestCase):
    def setUp(self):
        wb = Workbook()
        make_sheet(wb.active)
        wb.active.title = "RM Portugal AllUpdate"
        fd, self.path = tempfile.mkstemp(suffix=".xlsx")
        os.close(fd)
        wb.save(self.path)

    def tearDown(self):
        os.remove(self.path)

    def test_load(self):
        conf = [{"group": "B", "label": "Projeto B", "area_match": "Zona B", "name_match": "A", "project": "P"},
                {"group": "A", "label": "T1", "area_match": "Zona A", "name_match": "Torre I"},
                {"group": "A", "label": "T2", "area_match": "Zona A", "name_match": "Torre II"},
                {"label": "Domitys", "area_match": "Vila", "name_match": "Vila"},
                {"label": "Falta", "name_match": "Nao existe"}]
        rows, unmatched = roadmap.load(self.path, "RM Portugal AllUpdate", conf)
        self.assertEqual([r["label"] for r in rows], ["Projeto B", "T1", "T2", "Domitys"])  # ordem do config
        self.assertEqual(unmatched, ["Falta"])
        b, t1, t2, vila = rows
        self.assertEqual((b["pspa"], b["acquisition"], b["launch"]), ("2023-11-01", "2024-03-01", "2025-02-01"))
        self.assertEqual((b["construction_start"], b["construction_end"], b["end_deliveries"]),
                         ("2025-07-01", "2027-03-01", "2027-09-01"))
        self.assertEqual((b["project"], b["group"], b["apartments"]), ("P", "B", 80))
        self.assertEqual(t1["name"], "Torre I")      # nome exato: não apanha "Torre II"
        self.assertEqual(t2["area"], "Zona A")       # zona herdada da linha anterior
        self.assertEqual(vila["name"], "Vila")       # sem nome próprio: usa a zona

    def test_missing_sheet(self):
        with self.assertRaises(ValueError):
            roadmap.load(self.path, "Outra", [])


if __name__ == "__main__":
    unittest.main()
